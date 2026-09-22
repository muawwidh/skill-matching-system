from dataclasses import dataclass
from pathlib import Path

from fastapi import HTTPException, status

from app.services.pdf_text_normalization_service import PdfTextNormalizationService


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    content_type: str


class DocumentParserService:
    MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024
    SUPPORTED_CONTENT_TYPES = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text/plain",
    }

    def __init__(self) -> None:
        self.pdf_normalizer = PdfTextNormalizationService()

    def validate_upload(self, filename: str, content_type: str, content: bytes) -> None:
        if content_type not in self.SUPPORTED_CONTENT_TYPES:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail="Only PDF, DOCX, and TXT files are supported.",
            )
        if len(content) > self.MAX_FILE_SIZE_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="File is larger than the 5 MB limit.",
            )
        if not Path(filename).suffix:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file must include a file extension.",
            )

    def parse_upload(self, filename: str, content_type: str, content: bytes) -> ParsedDocument:
        self.validate_upload(filename, content_type, content)
        if content_type == "text/plain":
            return ParsedDocument(text=self._parse_txt(content), content_type=content_type)
        if content_type == "application/pdf":
            return ParsedDocument(text=self._parse_pdf(content), content_type=content_type)
        return ParsedDocument(text=self._parse_docx(content), content_type=content_type)

    def parse_text(self, text: str) -> ParsedDocument:
        stripped_text = text.strip()
        if not stripped_text:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="CV text is empty.")
        return ParsedDocument(text=stripped_text, content_type="text/plain")

    def _parse_txt(self, content: bytes) -> str:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        if not text.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="TXT file is empty.")
        return text

    def _parse_pdf(self, content: bytes) -> str:
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Scanned PDFs are not supported because no readable text was found.",
            )
        return self.pdf_normalizer.normalize(text)

    def _parse_docx(self, content: bytes) -> str:
        from docx import Document
        from io import BytesIO

        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if not text.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="DOCX file is empty.")
        return text
